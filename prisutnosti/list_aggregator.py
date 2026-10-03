"""Combine only lists downloaded in the current calendar run."""
import logging
from pathlib import Path
import tempfile

import pandas as pd
from openpyxl import load_workbook

logger = logging.getLogger(__name__)
METADATA_COLUMNS = ["TIP NASTAVE", "OD", "DO", "POČETAK PRIJAVE"]


class AttendanceListAggregator:
    def __init__(self, output_path: Path) -> None:
        self.output_path = output_path
        self.frames: list[pd.DataFrame] = []
        self.downloaded_files = 0
        self.output_conflict = False

    def add(self, filepath: Path, metadata: dict) -> None:
        if filepath.resolve() == self.output_path.resolve():
            self.output_conflict = True
            raise ValueError("Aggregate output cannot replace an individual downloaded list")
        # DataTables exports may place a merged report title above the headers.
        workbook = load_workbook(filepath, read_only=False, data_only=True)
        try:
            sheet = workbook.active
            title = any(area.min_row == area.max_row == 1 and area.max_col > 1
                        for area in sheet.merged_cells.ranges)
        finally:
            workbook.close()
        frame = pd.read_excel(filepath, header=1 if title else 0, dtype=object)
        collisions = set(METADATA_COLUMNS).intersection(frame.columns)
        if collisions:
            raise ValueError(f"Downloaded list already contains calendar columns: {sorted(collisions)}")
        for column in METADATA_COLUMNS:
            frame[column] = metadata[column]
        self.frames.append(frame)
        self.downloaded_files += 1
        logger.info("Adding %d student records from %s to aggregate", len(frame), filepath)

    def write(self) -> None:
        if self.output_conflict:
            raise ValueError("Aggregate output conflicts with an individual downloaded list")
        result = pd.concat(self.frames, ignore_index=True, sort=False) if self.frames else pd.DataFrame(columns=METADATA_COLUMNS)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(dir=self.output_path.parent, suffix=".xlsx", delete=False) as handle:
                temporary = Path(handle.name)
            result.to_excel(temporary, index=False, sheet_name="Prisutnosti")
            temporary.replace(self.output_path)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
        logger.info("Aggregate saved: %s; lists=%d student records=%d",
                    self.output_path, self.downloaded_files, len(result))
